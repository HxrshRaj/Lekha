import json
import shutil
import subprocess
from pathlib import Path

import pandas as pd
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook


@pytest.fixture(scope="module")
def client(pipeline_out):
    import lekha.api as api
    api.RESULTS = pipeline_out
    api.results.cache_clear()
    return TestClient(api.app)


def test_endpoints_ok(client):
    for u in ["/api/health", "/api/meta", "/api/kpis", "/api/aging", "/api/dso", "/api/otif", "/api/cycle-time", "/api/variance",
              "/api/pnl", "/api/forecast", "/api/regression", "/api/simulation", "/api/nlp/metrics", "/api/anomalies",
              "/api/anomalies?kind=payments", "/api/data-quality"]:
        assert client.get(u).status_code == 200, u
    assert "SYNTHETIC" in client.get("/api/health").json()["disclaimer"]
    assert "SYNTHETIC" in client.get("/openapi.json").json()["info"]["description"]


def test_validation_and_classify(client):
    assert client.get("/api/anomalies?kind=bogus").status_code == 422
    assert client.post("/api/nlp/classify", json={"text": "x"}).status_code == 422
    r = client.post("/api/nlp/classify", json={"text": "the delivery arrived damaged and the items were broken"})
    assert r.status_code == 200 and r.json()["label"] == "damaged_goods"
    assert abs(sum(s["probability"] for s in r.json()["scores"]) - 1) < 1e-6


def test_kpis_endpoint_matches_pipeline(client, pipeline_out):
    res = json.loads((pipeline_out / "results.json").read_text())
    assert client.get("/api/kpis").json() == res["kpis"]
    assert client.get("/api/variance?period=202512").json()["period"] == 202512


def test_export_xlsx_downloads(client):
    r = client.get("/export.xlsx")
    assert r.status_code == 200 and r.content[:2] == b"PK"


def test_excel_has_real_formulas_and_conditional_formatting(pipeline_out):
    wb = load_workbook(pipeline_out / "lekha_export.xlsx")
    assert {"Summary", "AR_Open", "Aging", "DSO", "Variance", "PnL"} <= set(wb.sheetnames)
    assert wb["AR_Open"]["G2"].value.startswith("=IF(")
    assert wb["Aging"]["C2"].value.startswith("=SUMIFS(")
    assert wb["DSO"]["D2"].value.startswith("=ROUND(")
    assert wb["Variance"]["E2"].value == "=C2-D2"
    assert len(wb["Aging"].conditional_formatting) > 0 and len(wb["Variance"].conditional_formatting) > 0
    assert "SYNTHETIC" in wb["Summary"]["A2"].value


def _soffice():
    return shutil.which("soffice") or shutil.which("libreoffice") or next(
        (p for p in ["C:/Program Files/LibreOffice/program/soffice.exe"] if Path(p).exists()), None)


@pytest.mark.skipif(_soffice() is None, reason="LibreOffice not installed in this environment")
def test_libreoffice_recalculation_matches_pipeline(pipeline_out, tmp_path):
    src = tmp_path / "in.xlsx"
    shutil.copy(pipeline_out / "lekha_export.xlsx", src)
    out = tmp_path / "out"
    subprocess.run([_soffice(), "--headless", "--calc", "--convert-to", "xlsx", "--outdir", str(out), str(src)],
                   check=True, timeout=240, capture_output=True)
    wb = load_workbook(out / "in.xlsx", data_only=True)
    res = json.loads((pipeline_out / "results.json").read_text())
    ag = {r["Bucket"]: r for r in res["ar_aging"]}
    for row in range(2, 7):
        b = wb["Aging"].cell(row, 1).value
        got = wb["Aging"].cell(row, 3).value or 0
        assert abs(got - ag.get(b, {"Amount": 0})["Amount"]) < 0.5, b
    assert abs(wb["Aging"]["C7"].value - res["kpis"]["open_ar"]) < 1
    assert abs(wb["Summary"]["B8"].value - res["kpis"]["dso_latest"]) < 0.01
    var = pd.DataFrame(res["variance"])
    last12 = var[var.PERIOD >= sorted(var.PERIOD.unique())[-12]]
    assert abs(wb["Summary"]["B9"].value - last12["Variance"].sum()) < 1
    assert abs(wb["Summary"]["B11"].value - res["kpis"]["ebit_ltm"]) < 1
